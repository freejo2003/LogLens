from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.test import TestCase
from docx import Document

from apps.ingestion.models import SourceFile
from apps.ingestion.services.docx_extractor import extract_docx_file
from apps.workspaces.models import Workspace


class DocxExtractorTests(TestCase):
    def setUp(self):
        User = get_user_model()

        self.user = User.objects.create_user(
            username="docx-test-user",
            password="test-password",
        )

        self.workspace = Workspace.objects.create(
            name="DOCX Test Workspace",
            owner=self.user,
        )

    def test_extracts_docx_paragraphs_into_segments(self):
        with TemporaryDirectory() as temp_dir:
            docx_path = Path(temp_dir) / "sample.docx"

            document = Document()
            document.add_paragraph("First paragraph")
            document.add_paragraph("Second paragraph")
            document.save(docx_path)

            source_file = SourceFile.objects.create(
                workspace=self.workspace,
                owner=self.user,
                original_name="sample.docx",
                storage_key="uploads/test-sample.docx",
                storage_path=str(docx_path),
                extension=".docx",
                detected_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                byte_size=docx_path.stat().st_size,
                checksum="test-checksum",
            )

            count = extract_docx_file(source_file)

            self.assertEqual(count, 2)

            segments = list(
                source_file.text_segments.order_by("sequence")
            )

            self.assertEqual(segments[0].text, "First paragraph")
            self.assertEqual(segments[0].paragraph_number, 1)
            self.assertEqual(segments[1].text, "Second paragraph")
            self.assertEqual(segments[1].paragraph_number, 2)

            self.assertEqual(
                segments[0].parser_name,
                "python-docx",
            )
            self.assertEqual(
                segments[0].parser_version,
                "1.0",
            )

    def test_rejects_non_docx_file(self):
        source_file = SourceFile(
            workspace=self.workspace,
            owner=self.user,
            original_name="test.txt",
            storage_key="uploads/test.txt",
            storage_path="missing.txt",
            extension=".txt",
            detected_type="text/plain",
            byte_size=10,
            checksum="test-checksum",
        )

        with self.assertRaises(ValueError):
            extract_docx_file(source_file)

    def test_raises_when_docx_does_not_exist(self):
        source_file = SourceFile(
            workspace=self.workspace,
            owner=self.user,
            original_name="missing.docx",
            storage_key="uploads/missing.docx",
            storage_path="missing.docx",
            extension=".docx",
            detected_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            byte_size=10,
            checksum="test-checksum",
        )

        with self.assertRaises(FileNotFoundError):
            extract_docx_file(source_file)