from docx import Document

from apps.ingestion.models import SourceFile, TextSegment
from apps.ingestion.services.errors import ExtractionError
from apps.ingestion.services.segment_writer import SegmentBatchWriter
from apps.ingestion.services.validators import _validate_docx_archive
from pathlib import Path


DOCX_PARSER_NAME = "python-docx"
DOCX_PARSER_VERSION = "1.1"


def extract_docx_file(source_file: SourceFile) -> int:
    if source_file.extension != ".docx":
        raise ExtractionError("This Word parser does not support the uploaded file type.")

    source_path = Path(source_file.storage_path)
    if not source_path.exists() or not source_path.is_file():
        raise ExtractionError("The original upload is unavailable. Upload the file again.")

    try:
        with source_path.open("rb") as file_obj:
            _validate_docx_archive(file_obj)
        document = Document(str(source_path))
    except Exception as exc:
        raise ExtractionError("The DOCX document could not be read. Verify that it is valid.") from exc

    TextSegment.objects.filter(source_file=source_file).delete()
    writer = SegmentBatchWriter(source_file)
    try:
        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            writer.add(
                paragraph.text,
                paragraph_number=paragraph_number,
                parser_name=DOCX_PARSER_NAME,
                parser_version=DOCX_PARSER_VERSION,
            )

        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                for cell_number, cell in enumerate(row.cells, start=1):
                    writer.add(
                        cell.text,
                        section_label=f"Table {table_number}, row {row_number}, cell {cell_number}",
                        parser_name=DOCX_PARSER_NAME,
                        parser_version=DOCX_PARSER_VERSION,
                    )
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("The DOCX text could not be extracted safely.") from exc

    if writer.character_count == 0:
        raise ExtractionError("The DOCX contains no extractable text.")
    return writer.finish()
